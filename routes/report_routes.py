#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Report Routes
Handles report generation and export endpoints
"""

import logging
import os
import uuid
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file
from src.services.report_service import ReportService
from middleware.rate_limiter import auto_rate_limit
from middleware.auth_middleware import require_auth

logger = logging.getLogger(__name__)

report_bp = Blueprint('report', __name__)

# Initialize report service
report_service = ReportService("Route: report_routes global")


@report_bp.route('/api/report', methods=['POST'])
@require_auth
@auto_rate_limit
def generate_report():
    """API สำหรับสร้างรายงาน"""
    try:
        data = request.get_json()
        report_date = data.get('report_date')
        job_type_id = data.get('job_type_id')
        sub_job_type_id = data.get('sub_job_type_id')
        note_filter = data.get('note_filter')
        barcode_filter = data.get('barcode_filter')

        # report_date is now optional (None = all dates)

        # Convert IDs to integers (job_type_id is now optional)
        try:
            job_type_id = int(job_type_id) if job_type_id else None
            sub_job_type_id = int(sub_job_type_id) if sub_job_type_id else None
        except ValueError:
            return jsonify({
                'success': False,
                'message': 'ID ไม่ถูกต้อง'
            })

        # Generate report using ReportService
        result = report_service.generate_report(
            report_date=report_date,
            job_id=job_type_id,
            sub_job_id=sub_job_type_id,
            note_filter=note_filter,
            barcode_filter=barcode_filter
        )
        
        if result['success']:
            # Convert datetime objects to strings for JSON serialization
            for record in result['data']:
                if hasattr(record['scan_date'], 'isoformat'):
                    record['scan_date'] = record['scan_date'].isoformat()
        
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"❌ เกิดข้อผิดพลาดในการสร้างรายงาน: {str(e)}")
        return jsonify({
            'success': False, 
            'message': f'เกิดข้อผิดพลาดในการสร้างรายงาน: {str(e)}'
        })


@report_bp.route('/api/report/export', methods=['POST'])
@require_auth
@auto_rate_limit
def export_report():
    """API สำหรับส่งออกรายงานเป็น Excel แล้วส่ง download URL กลับ"""
    try:
        data = request.get_json()
        report_data = data.get('report_data', [])
        summary = data.get('summary', {})

        if not report_data:
            return jsonify({
                'success': False,
                'message': 'ไม่มีข้อมูลสำหรับส่งออก'
            })

        # Generate Excel file
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            return jsonify({
                'success': False,
                'message': 'ไม่พบ openpyxl กรุณาติดตั้ง: pip install openpyxl'
            })

        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = "ข้อมูล"
        columns = [
            ("บาร์โค้ด", "barcode", 25),
            ("วันที่/เวลา", "scan_date", 22),
            ("งานหลัก", "job_type_name", 25),
            ("งานรอง", "sub_job_type_name", 25),
            ("หมายเหตุ", "notes", 30),
            ("ผู้ใช้", "user_id", 15),
        ]

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font_white = Font(bold=True, color="FFFFFF", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        wrap_align = Alignment(wrap_text=True, vertical="top")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Write headers
        for col_idx, (header, _, width) in enumerate(columns, 1):
            cell = ws_data.cell(row=1, column=col_idx, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws_data.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        # Date number format for Excel
        date_number_format = 'DD/MM/YYYY HH:MM:SS'

        # Write data rows
        for row_idx, item in enumerate(report_data, 2):
            for col_idx, (_, key, _) in enumerate(columns, 1):
                value = item.get(key, '')
                if key == 'scan_date' and value:
                    # Parse ISO date string to datetime for proper Excel formatting
                    try:
                        dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
                        cell = ws_data.cell(row=row_idx, column=col_idx, value=dt)
                        cell.number_format = date_number_format
                    except (ValueError, TypeError):
                        cell = ws_data.cell(row=row_idx, column=col_idx, value=str(value))
                else:
                    cell = ws_data.cell(row=row_idx, column=col_idx, value=str(value) if value else '')
                cell.border = thin_border
                if key == 'notes':
                    cell.alignment = wrap_align

        # Save to temp file
        export_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'temp_exports')
        os.makedirs(export_dir, exist_ok=True)

        file_id = uuid.uuid4().hex[:8]
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"report_{timestamp}_{file_id}.xlsx"
        filepath = os.path.join(export_dir, filename)
        wb.save(filepath)

        return jsonify({
            'success': True,
            'message': 'ส่งออกสำเร็จ',
            'download_url': f'/api/report/download/{filename}'
        })

    except Exception as e:
        logger.error(f"Error exporting report: {str(e)}")
        return jsonify({
            'success': False,
            'message': f'เกิดข้อผิดพลาดในการส่งออก: {str(e)}'
        })


@report_bp.route('/api/report/download/<filename>')
@require_auth
def download_report(filename):
    """ดาวน์โหลดไฟล์ Excel ที่ส่งออก"""
    try:
        # Validate filename to prevent path traversal
        if '..' in filename or '/' in filename or '\\' in filename:
            return jsonify({'success': False, 'message': 'ชื่อไฟล์ไม่ถูกต้อง'}), 400

        export_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'temp_exports')
        filepath = os.path.join(export_dir, filename)

        if not os.path.exists(filepath):
            return jsonify({'success': False, 'message': 'ไม่พบไฟล์'}), 404

        return send_file(
            filepath,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        logger.error(f"Error downloading report: {str(e)}")
        return jsonify({'success': False, 'message': str(e)}), 500


@report_bp.route('/api/report/monthly/<int:year>/<int:month>')
@auto_rate_limit
def get_monthly_summary(year, month):
    """API สำหรับดึงสรุปรายเดือน"""
    try:
        # Validate year and month
        if not (1900 <= year <= 2100) or not (1 <= month <= 12):
            return jsonify({
                'success': False,
                'message': 'ปีหรือเดือนไม่ถูกต้อง'
            })
        
        result = report_service.get_monthly_summary(year, month)
        
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"Error getting monthly summary: {str(e)}")
        return jsonify({
            'success': False,
            'message': f'เกิดข้อผิดพลาด: {str(e)}'
        })


@report_bp.route('/api/report/user_activity')
@auto_rate_limit
def get_user_activity_report():
    """API สำหรับดึงรายงานกิจกรรมผู้ใช้"""
    try:
        start_date = request.args.get('start_date')
        end_date = request.args.get('end_date')
        
        if not start_date or not end_date:
            return jsonify({
                'success': False,
                'message': 'กรุณาระบุวันที่เริ่มต้นและสิ้นสุด'
            })
        
        result = report_service.get_user_activity_report(start_date, end_date)
        
        # Convert datetime objects to strings for JSON serialization
        if result['success']:
            for activity in result['data']['user_activities']:
                if hasattr(activity['first_scan'], 'isoformat'):
                    activity['first_scan'] = activity['first_scan'].isoformat()
                if hasattr(activity['last_scan'], 'isoformat'):
                    activity['last_scan'] = activity['last_scan'].isoformat()
        
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"Error getting user activity report: {str(e)}")
        return jsonify({
            'success': False,
            'message': f'เกิดข้อผิดพลาด: {str(e)}'
        })